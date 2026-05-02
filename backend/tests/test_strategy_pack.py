from pathlib import Path
import sys


BACKEND_DIR = Path(__file__).resolve().parents[1]
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from app.services.export_engine import PPTXGenerator  # noqa: E402
from app.services.strategy_pack import StrategyPackService  # noqa: E402


def _payload(source_mode="demo_mode"):
    return {
        "mode": "agency_client_pitch_summary",
        "metadata": {
            "agency_name": "Demo Agency",
            "client_name": "Sample Client",
            "prepared_by": "Strategy Team",
            "report_date": "2026-05-02",
            "campaign_name": "Premium Water Launch",
            "scenario_name": "Launch readiness",
            "logo_url": "https://example.test/logo.png",
        },
        "campaign": {
            "campaign_id": "cmp_demo",
            "name": "Premium Water Launch",
            "objective": "launch readiness",
        },
        "kpis": {
            "overall_sentiment": 28,
            "conversion_probability": 62,
            "message_resonance": 71,
            "crisis_risk": 42,
            "brand_perception_shift": 8,
            "opinion_polarization": 31,
            "confidence_score": 64,
        },
        "segments": [
            {"name": "Urban premium buyers", "sentiment": 48},
            {"name": "Price-sensitive families", "sentiment": -22},
        ],
        "evidence": {
            "risk_drivers": ["Price premium needs clearer proof."],
            "crisis_watchouts": ["Monitor price backlash comments."],
            "assumptions": ["Synthetic sample for product walkthrough."],
            "limitations": ["Requires audience validation before launch."],
        },
        "action_plan": {
            "sections": {
                "creative_adjustment": {
                    "title": "Creative Adjustment",
                    "items": [
                        {
                            "recommendation": "Add proof cards to the hero claim.",
                            "reason": "Premium positioning needs clearer support.",
                            "expected_impact": "Improves message clarity.",
                            "risk": "May reduce emotional simplicity.",
                        }
                    ],
                }
            }
        },
        "source": {
            "type": source_mode,
            "source_mode": source_mode,
            "data_basis": {
                "demo_mode": "demo_fixture",
                "local_estimate": "local_estimate",
            }.get(source_mode, "real_simulation"),
            "run_id": "run_demo",
        },
    }


def test_strategy_pack_has_client_ready_sections_and_provenance():
    pack = StrategyPackService().build(_payload())

    assert pack["version"] == "strategy_pack_v1"
    assert pack["mode"] == "agency_client_pitch_summary"
    assert pack["mode_label"] == "Agency Client Pitch Summary"
    assert pack["source"]["source_mode"] == "demo_mode"
    assert pack["source"]["data_basis"] == "demo_fixture"
    assert pack["disclaimer"]

    for key in pack["section_order"]:
        section = pack["sections"][key]
        assert section["title"]
        assert "items" in section
        assert section["provenance"]["source_mode"] == "demo_mode"
        assert section["provenance"]["data_basis"] == "demo_fixture"
        assert section["provenance"]["confidence_level"] == "medium"
        assert section["provenance"]["assumptions"]
        assert section["provenance"]["limitations"]
        assert section["provenance"]["recommended_next_validation_step"]


def test_strategy_pack_distinguishes_brand_and_agency_modes():
    service = StrategyPackService()
    brand = service.build({**_payload("backend_verified"), "mode": "brand"})
    agency = service.build({**_payload("backend_verified"), "mode": "agency"})

    assert brand["mode"] == "brand_executive_summary"
    assert brand["mode_label"] == "Brand Executive Summary"
    assert brand["audience"] == "brand_leadership"
    assert agency["mode"] == "agency_client_pitch_summary"
    assert agency["mode_label"] == "Agency Client Pitch Summary"
    assert agency["audience"] == "agency_client_team"
    assert brand["source"]["source_mode"] == "backend_verified"


def test_strategy_pack_does_not_echo_logo_url_or_upgrade_source_labels():
    payload = _payload("local_estimate")
    pack = StrategyPackService().build(payload)

    assert pack["white_label"]["logo_url_present"] is True
    assert "logo_url" not in pack["white_label"]
    assert "example.test/logo" not in str(pack)
    assert pack["source"]["source_mode"] == "local_estimate"
    assert "backend_verified" not in str(pack["source"])


def _slide_text(slide):
    parts = []
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False):
            parts.append(shape.text)
    return "\n".join(parts)


def test_strategy_pack_pptx_renderer_preserves_source_metadata_on_every_slide():
    from pptx import Presentation

    pack = StrategyPackService().build(_payload("demo_mode"))
    buffer = PPTXGenerator().generate({
        "slide_type": "strategy_pack",
        "strategy_pack": pack,
    })

    prs = Presentation(buffer)
    slide_texts = [_slide_text(slide) for slide in prs.slides]
    deck_text = "\n".join(slide_texts)

    assert len(prs.slides) == 7
    assert "Agency Client Pitch Summary" in deck_text
    assert "Demo Agency" in deck_text
    assert "Sample Client" in deck_text
    assert "Limitations & Recommended Validation" in deck_text
    assert "Recommended validation:" in deck_text
    assert "Decision-support scenario planning" in deck_text
    assert "guaranteed prediction" in deck_text
    assert "example.test/logo" not in deck_text

    for text in slide_texts:
        assert "Source: demo_mode" in text
        assert "Basis: demo_fixture" in text
        assert "Run ID: run_demo" in text


def test_strategy_pack_pptx_renderer_supports_brand_and_agency_modes():
    from pptx import Presentation

    for mode, expected_label in [
        ("brand", "Brand Executive Summary"),
        ("agency", "Agency Client Pitch Summary"),
    ]:
        pack = StrategyPackService().build({**_payload("local_estimate"), "mode": mode})
        buffer = PPTXGenerator().generate({
            "slide_type": "strategy_pack",
            "strategy_pack": pack,
        })
        text = "\n".join(_slide_text(slide) for slide in Presentation(buffer).slides)

        assert expected_label in text
        assert "Source: local_estimate" in text
        assert "Basis: local_estimate" in text
